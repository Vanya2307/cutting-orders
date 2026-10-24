from openpyxl import Workbook

from django.utils import timezone
from openpyxl.styles import Alignment, Font, Border, Side

def generate_order_excel(order):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Поръчка'

    border_side = Side(style='thin', color='000000')
    cell_border = Border(
        left=border_side,
        right=border_side,
        top=border_side,
        bottom=border_side,
    )

    headers = [
        'Ref №',
        'Материал',
        'D Фладер',
        'Наименование',
        'L Дължина',
        'W Ширина',
        'Бр.',
        'Кантиране',
        '',
        '',
        '',
        'Друго',
        'Дебелина мм',
        'Клиент',
        'Обработка',
    ]

    for column, heading in enumerate(headers, start=1):
        worksheet.cell(row=2, column=column, value=heading)

    edge_headers = ['L горе', 'L долу', 'W ляво', 'W дясно']

    for column, heading in enumerate(edge_headers, start=8):
        worksheet.cell(row=3, column=column, value=heading)


    submitted_date = (
        timezone.localtime(order.submitted_at).strftime('%d.%m.%Y')
        if order.submitted_at
        else ''
    )

    worksheet['A1'] = f'Пор. №: {order.number or ""}'
    worksheet['C1'] = f'Име: {order.contact_name}'
    worksheet['G1'] = f'Тел: {order.contact_phone}'
    worksheet['L1'] = f'Дата: {submitted_date}'

    # merging the cells at first row
    for cell_range in ('A1:B1', 'C1:F1', 'G1:K1', 'L1:O1'):
        worksheet.merge_cells(cell_range)

    worksheet.merge_cells('H2:K2')

    for column in ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'L', 'M', 'N', 'O'):
        worksheet.merge_cells(f'{column}2:{column}3')

    # Widths of column

    column_widths = {
        'A': 6,
        'B': 20,
        'C': 8,
        'D': 20,
        'E': 10,
        'F': 10,
        'G': 6,
        'H': 9,
        'I': 9,
        'J': 9,
        'K': 9,
        'L': 22,
        'M': 12,
        'N': 20,
        'O': 28,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # Aligning order information
    for address in ('A1', 'C1', 'G1', 'L1'):
        worksheet[address].font = Font(bold=True)
        worksheet[address].alignment = Alignment(vertical='center', wrap_text=True)

    # Alignment of headings
    for row in worksheet.iter_rows(min_row=2, max_row=3, max_col=15):
        for cell in row:
            cell.border = cell_border
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal='center',
                vertical='center',
                wrap_text=True,
            )

    worksheet.row_dimensions[1].height = 36
    worksheet.row_dimensions[2].height = 30
    worksheet.row_dimensions[3].height = 24


    details = (
        order.details
        .select_related('material')
        .prefetch_related('services')
        .order_by('material__decor_code', 'material_id', 'pk')
    )

    row_number = 4
    previous_material_id = None
    reference_number = 0

    for detail in details:
        material = detail.material

        # 2 empty rows between different materials

        if (previous_material_id is not None and previous_material_id != material.pk):
            row_number += 2
            reference_number = 0

        reference_number += 1

        processing = detail.processing.strip() or ', '.join(
            service.name for service in detail.services.all()
        )

        processing_lines = [processing] if processing else []

        if detail.euro_joint_left != 'NO_JOINT':
            processing_lines.append(
                'Евросглобка, лява: '
                f'{detail.get_euro_joint_left_display()}'
            )

        if detail.euro_joint_right != 'NO_JOINT':
            processing_lines.append(
                'Евросглобка, дясна: '
                f'{detail.get_euro_joint_right_display()}'
            )

        grain = (
            detail.grain_direction
            if material.has_grain and detail.grain_direction in ('L', 'W')
            else None
        )

        values = [
            reference_number,
            material.get_label(),
            grain,
            detail.name,
            detail.length,
            detail.width,
            detail.quantity,
            float(detail.edge_l_top) if detail.edge_l_top else None,
            float(detail.edge_l_bottom) if detail.edge_l_bottom else None,
            float(detail.edge_w_left) if detail.edge_w_left else None,
            float(detail.edge_w_right) if detail.edge_w_right else None,
            detail.other,
            float(material.thickness),
            f'{order.number}-{order.contact_name}',
            '\n'.join(processing_lines),

        ]


        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(
                row=row_number,
                column=column,
                value=value,
            )

            cell.border = cell_border

            if isinstance(value, str):
                cell.data_type = 's'

            cell.alignment = Alignment(
                horizontal='left' if column == 2 else 'center',
                vertical='center',
                wrap_text=column in (2, 12, 14, 15),
            )

        previous_material_id = material.pk
        row_number += 1


    for cell in worksheet[1]:
        cell.border = Border(
            bottom=Side(style='medium', color='000000'),
        )

    return workbook





